import calendar as _calendar
import io
import json
import os
import secrets
import zipfile
from datetime import date, datetime, timedelta

from django.utils.timezone import now as timezone_now
from django.contrib.auth.decorators import login_required, permission_required
from django.db.models import Count, F, IntegerField, OuterRef, Q, Subquery, Value
from django.db.models.functions import Coalesce
from django.http import FileResponse, HttpResponse, HttpResponseBadRequest
from django.shortcuts import get_object_or_404, redirect, render
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.decorators.http import require_POST
from django_tables2 import RequestConfig
from ._pagination import ist_nachladen, nachlade_antwort, paginate_ordered, parse_offset
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from icalendar import Calendar, Event

from intern import folders
from intern.app_settings import get_config

from ..forms import TerminForm
from ..models import Konzert, Lied, Ort, Termin, User, Zusage
from ..tables import TerminTabelle

__all__ = [
    "termin_list",
    "termin_kalender",
    "termin_view",
    "termin_zusage",
    "termin_anwesenheit_update",
    "termin_edit",
    "termin_delete",
    "termin_add",
    "calendar_ics",
    "konzert_docx",
    "konzert_docx_titles",
    "konzert_pdf",
    "konzert_pdf_zip",
    "konzert_noten_zip",
]

ZUSAGE_TRANSITIONS = {
    "": ["1", "2", "3"],
    "3": ["1", "2"],
    "1": [],
    "2": ["1", "3"],
}


def _build_stimme_groups(termin):
    STIMMEN_ORDER = ["S", "A", "T", "B", "C", "M"]
    STIMMEN_LABELS = {"S": "Sopran", "A": "Alt", "T": "Tenor", "B": "Bass", "C": "Chorleiter", "M": "Musiker"}
    active_users = User.objects.filter(is_active=True).order_by("last_name", "first_name")
    zusagen = {z.person_id: z.option for z in Zusage.objects.filter(termin=termin)}
    groups = []
    for stimme in STIMMEN_ORDER:
        users = [u for u in active_users if u.stimme == stimme]
        if not users:
            continue
        g = {"1": [], "2": [], "3": [], "": []}
        for u in users:
            g[zusagen.get(u.id, "")].append(u)
        groups.append({
            "key": stimme, "label": STIMMEN_LABELS[stimme],
            "zusagen": g["1"], "absagen": g["2"],
            "vielleicht": g["3"], "keine_angabe": g[""],
        })
    return groups


def calendar_ics(request, secret):
    ical_secret = get_config("ICAL_SECRET") or ""
    if not ical_secret or not secrets.compare_digest(secret, ical_secret):
        from django.http import Http404
        raise Http404
    cutoff = timezone_now() - timedelta(days=180)
    termine = Termin.objects.filter(datum__gte=cutoff).select_related("ort").order_by("datum")
    cal = Calendar()
    cal.add("prodid", "-//CantorHub//DE")
    cal.add("version", "2.0")
    cal.add("x-wr-calname", "CantorHub")
    for t in termine:
        ev = Event()
        ev.add("uid", f"termin-{t.pk}@cantorhub")
        ev.add("summary", t.kurz)
        ev.add("dtstart", t.datum)
        ev.add("dtend", t.datum_bis or t.datum + timedelta(hours=2))
        if t.lang:
            ev.add("description", t.lang)
        if t.ort:
            parts = [t.ort.name, t.ort.strasse, f"{t.ort.plz or ''} {t.ort.ort or ''}".strip()]
            ev.add("location", ", ".join(p for p in parts if p))
            if t.ort.breite is not None and t.ort.laenge is not None:
                ev.add("geo", (t.ort.breite, t.ort.laenge))
        cal.add_component(ev)
    return HttpResponse(cal.to_ical(), content_type="text/calendar; charset=utf-8")


@login_required
def termin_list(request):
    past = request.GET.get("past") == "1"
    annotations = {
        "anwesenheit_option": Coalesce(
            Subquery(
                Zusage.objects.filter(
                    termin=OuterRef("pk"), person=request.user
                ).values("option")[:1]
            ),
            Value(""),
        ),
        "anzahl_songs": Count("konzert", distinct=True),
        "anzahl_zusagen_sopran": Count(
            "zusage",
            filter=Q(zusage__person__stimme="S", zusage__option="1"),
            distinct=True,
        ),
        "anzahl_zusagen_alt": Count(
            "zusage",
            filter=Q(zusage__person__stimme="A", zusage__option="1"),
            distinct=True,
        ),
        "anzahl_zusagen_tenor": Count(
            "zusage",
            filter=Q(zusage__person__stimme="T", zusage__option="1"),
            distinct=True,
        ),
        "anzahl_zusagen_bass": Count(
            "zusage",
            filter=Q(zusage__person__stimme="B", zusage__option="1"),
            distinct=True,
        ),
        "anzahl_vielleicht": Count(
            "zusage",
            filter=Q(zusage__option="3"),
            distinct=True,
        ),
        "anzahl_absagen": Count(
            "zusage",
            filter=Q(zusage__option="2"),
            distinct=True,
        ),
        "anzahl_keine_angabe": Coalesce(
            Value(
                User.objects.filter(is_active=True).count(),
                output_field=IntegerField(),
            ) - Count("zusage", distinct=True),
            Value(0),
            output_field=IntegerField(),
        ),
    }
    if past:
        qs = Termin.objects.filter(datum__lte=timezone_now()).order_by("-datum")
    else:
        qs = Termin.objects.filter(datum__gt=timezone_now()).order_by("datum")
    termin_queryset = qs.annotate(**annotations)

    _PAGE_SIZE = 50
    offset = parse_offset(request)

    probe = TerminTabelle(termin_queryset)
    rows, has_more = paginate_ordered(request, probe, offset, _PAGE_SIZE)
    table = TerminTabelle(rows)
    RequestConfig(request, paginate=False).configure(table)

    if ist_nachladen(request):
        return nachlade_antwort("table_rows.html", {"table": table},
                                has_more, offset + _PAGE_SIZE)

    from django.urls import reverse

    ical_secret = get_config("ICAL_SECRET")
    if ical_secret:
        ical_path = reverse("calendar_ics", args=[ical_secret])
        ical_url = request.build_absolute_uri(ical_path)
        webcal_url = "webcal://" + request.get_host() + ical_path
    else:
        ical_url = webcal_url = None
    add = "termin_add" if request.user.has_perm("intern.add_termin") else None
    return render(
        request,
        "table.html",
        {
            "table": table,
            "seitentitel": "Vergangene Termine" if past else "Termine",
            "topbar_icon": '<svg class="topbar-icon" xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><rect x="3" y="4" width="18" height="18" rx="2" ry="2"/><line x1="16" y1="2" x2="16" y2="6"/><line x1="8" y1="2" x2="8" y2="6"/><line x1="3" y1="10" x2="21" y2="10"/></svg>',
            "ical_url": ical_url,
            "webcal_url": webcal_url,
            "upload_folder_id": folders.termine(),
            "upload_folder_name": "Termine",
            "add": add,
            "add_form": TerminForm() if request.user.has_perm("intern.add_termin") else None,
            "add_label": "Neuer Termin",
            "past": past,
            "has_more": has_more,
            "next_offset": offset + _PAGE_SIZE,
        },
    )


@login_required
def termin_view(request, id):
    termin = get_object_or_404(Termin.objects.select_related("ort"), id=id)
    lieder = Konzert.objects.filter(termin=termin).select_related("lied").order_by("position")
    my_zusage = Zusage.objects.filter(termin=termin, person=request.user).order_by("-id").first()
    my_option = my_zusage.option if my_zusage else ""
    context = {
        "seitentitel": str(termin),
        "termin": termin,
        "stimme_groups": _build_stimme_groups(termin),
        "my_option": my_option,
        "lieder": lieder,
        "up": "termin_list",
        "edit": "termin_edit" if request.user.has_perm("intern.change_termin") else None,
        "history": "termin_aenderungen" if request.user.has_perm("intern.change_termin") else None,
        "id": id,
    }
    return render(request, "termin.html", context)


@login_required
def konzert_docx(request, id):
    termin = get_object_or_404(Termin, id=id)
    lieder = Konzert.objects.filter(termin=termin).select_related("lied").order_by("position")

    doc = Document()

    titel = doc.add_paragraph(str(termin))
    titel.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titel.runs[0]
    run.bold = True
    run.font.size = Pt(18)

    for k in lieder:
        if k.lied:
            item = doc.add_paragraph(k.lied.titel)
            item.runs[0].font.size = Pt(14)
            if k.lied.komponist:
                comp = doc.add_paragraph(k.lied.komponist)
                comp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                comp.runs[0].font.size = Pt(10)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    safe_name = termin.kurz.replace('"', "")
    return FileResponse(buf, as_attachment=True, filename=f"{safe_name}.docx", content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


@login_required
def konzert_docx_titles(request, id):
    termin = get_object_or_404(Termin, id=id)
    lieder = Konzert.objects.filter(termin=termin).select_related("lied").order_by("position")

    doc = Document()

    titel = doc.add_paragraph(str(termin))
    titel.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = titel.runs[0]
    run.bold = True
    run.font.size = Pt(18)

    for k in lieder:
        if k.lied:
            item = doc.add_paragraph(k.lied.titel)
            item.alignment = WD_ALIGN_PARAGRAPH.CENTER
            item.runs[0].font.size = Pt(14)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    safe_name = termin.kurz.replace('"', "")
    return FileResponse(buf, as_attachment=True, filename=f"{safe_name}-titel.docx", content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


@login_required
def konzert_pdf(request, id):
    import fitz

    termin = get_object_or_404(Termin, id=id)
    lieder = Konzert.objects.filter(termin=termin).select_related("lied").order_by("position")

    combined = fitz.open()
    for k in lieder:
        if k.lied and k.lied.pdf:
            src = fitz.open(k.lied.pdf.file.path)
            combined.insert_pdf(src)
            src.close()

    buf = io.BytesIO()
    combined.save(buf)
    combined.close()
    buf.seek(0)

    safe_name = termin.kurz.replace('"', "")
    return FileResponse(buf, as_attachment=True, filename=f"{safe_name}.pdf", content_type="application/pdf")


@login_required
def konzert_pdf_zip(request, id):
    termin = get_object_or_404(Termin, id=id)
    lieder = Konzert.objects.filter(termin=termin).select_related("lied").order_by("position")

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for k in lieder:
            if k.lied and k.lied.pdf:
                arcname = f"{k.position:02d} {k.lied.titel}.pdf"
                zf.write(k.lied.pdf.file.path, arcname)
    buf.seek(0)

    safe_name = termin.kurz.replace('"', "")
    return FileResponse(buf, as_attachment=True, filename=f"{safe_name}-pdfs.zip", content_type="application/zip")


@login_required
def konzert_noten_zip(request, id):
    termin = get_object_or_404(Termin, id=id)
    lieder = Konzert.objects.filter(termin=termin).select_related("lied").order_by("position")

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for k in lieder:
            if k.lied and k.lied.noten:
                ext = os.path.splitext(k.lied.noten.file.path)[1]
                arcname = f"{k.position:02d} {k.lied.titel}{ext}"
                zf.write(k.lied.noten.file.path, arcname)
    buf.seek(0)

    safe_name = termin.kurz.replace('"', "")
    return FileResponse(buf, as_attachment=True, filename=f"{safe_name}-noten.zip", content_type="application/zip")


@login_required
@require_POST
def termin_zusage(request, id):
    from ..signals import _log_change
    from ..models import ChangeLog

    termin = get_object_or_404(Termin, id=id)
    if not termin.anwesenheit:
        if not request.headers.get("HX-Request"):
            return redirect("termin_view", id=id)
        return render(request, "_termin_zusage.html", {
            "termin": termin,
            "my_option": "",
            "stimme_groups": [],
        })
    new_option = request.POST.get("option", "")
    current = Zusage.objects.filter(termin=termin, person=request.user).order_by("-id").first()
    current_option = current.option if current else ""
    old_option = current_option

    if new_option in ZUSAGE_TRANSITIONS.get(current_option, []):
        obj, created = Zusage.objects.update_or_create(
            person=request.user, termin=termin,
            defaults={"option": new_option},
        )
        _log_change(
            obj, request.user,
            ChangeLog.ACTION_CREATE if created else ChangeLog.ACTION_UPDATE,
            old_values={"option": old_option},
        )
        current_option = new_option

    if not request.headers.get("HX-Request"):
        return redirect("termin_view", id=id)

    if request.POST.get("source") == "banner":
        return render(request, "_termin_banner.html", {
            "banner_termin": termin,
            "banner_my_option": current_option,
        })

    return render(request, "_termin_zusage.html", {
        "termin": termin,
        "my_option": current_option,
        "stimme_groups": _build_stimme_groups(termin),
    })


@permission_required("intern.change_zusage", raise_exception=True)
@require_POST
def termin_anwesenheit_update(request, id):
    from ..models import ChangeLog
    from ..signals import _log_change

    termin = get_object_or_404(Termin, id=id)
    if not termin.anwesenheit:
        return HttpResponseBadRequest("anwesenheit disabled")
    try:
        payload = json.loads(request.body.decode("utf-8") or "{}")
    except ValueError:
        return HttpResponseBadRequest("invalid json")

    changes = payload.get("changes") or []
    valid_options = {"1", "2", "3"}
    active_ids = set(User.objects.filter(is_active=True).values_list("id", flat=True))
    existing = {
        z.person_id: z for z in Zusage.objects.filter(termin=termin)
    }

    for entry in changes:
        try:
            person_id = int(entry.get("person_id"))
        except (TypeError, ValueError):
            continue
        option = entry.get("option")
        if option not in valid_options or person_id not in active_ids:
            continue
        current = existing.get(person_id)
        if current and current.option == option:
            continue
        old_option = current.option if current else None
        obj, created = Zusage.objects.update_or_create(
            person_id=person_id, termin=termin,
            defaults={"option": option},
        )
        _log_change(
            obj, request.user,
            ChangeLog.ACTION_CREATE if created else ChangeLog.ACTION_UPDATE,
            old_values={"option": old_option},
        )

    my_zusage = Zusage.objects.filter(termin=termin, person=request.user).order_by("-id").first()
    return render(request, "_termin_zusage.html", {
        "termin": termin,
        "my_option": my_zusage.option if my_zusage else "",
        "stimme_groups": _build_stimme_groups(termin),
    })


@permission_required("intern.change_termin", raise_exception=True)
def termin_edit(request, id):
    termin = get_object_or_404(Termin, id=id)

    if request.method == "POST":
        form = TerminForm(request.POST, instance=termin)
        if form.is_valid():
            instance = form.save(commit=False)
            instance._audit_user = request.user
            instance.save()
            form.save_m2m()
            return redirect("termin_view", id=id)
    else:
        form = TerminForm(instance=termin)

    # Was ein Loeschen mitnehmen wuerde — der Bestaetigungsdialog
    # (_confirm_delete_dialog.html) zeigt die Zeilen unveraendert an.
    warnungen = []
    zusagen = Zusage.objects.filter(termin=termin).count()
    if zusagen:
        warnungen.append(
            f"{zusagen} Zusage{'n werden' if zusagen != 1 else ' wird'} ebenfalls gelöscht."
        )
    konzerte = Konzert.objects.filter(termin=termin).count()
    if konzerte:
        warnungen.append(
            f"{konzerte} Konzert{'e werden' if konzerte != 1 else ' wird'} ebenfalls gelöscht."
        )

    context = {
        "form": form,
        "termin": termin,
        "seitentitel": f"bearbeiten: {termin.kurz}",
        "view": "termin_view",
        "id": id,
        "loesch_warnungen": warnungen,
        "can_delete": request.user.has_perm("intern.delete_termin"),
    }
    return render(request, "termin_edit.html", context)


@permission_required("intern.delete_termin", raise_exception=True)
@require_POST
def termin_delete(request, id):
    termin = get_object_or_404(Termin, id=id)
    termin._audit_user = request.user
    termin.delete()
    return redirect("termin_list")


@permission_required("intern.add_termin", raise_exception=True)
def termin_add(request):
    # Das Kalender-Offcanvas schickt ein "next" mit, damit man nach dem Anlegen
    # im angezeigten Monat landet statt auf der Detailseite. Nur interne Ziele.
    weiter = request.POST.get("next") or request.GET.get("next") or ""
    if weiter and not url_has_allowed_host_and_scheme(
        weiter, allowed_hosts={request.get_host()}, require_https=request.is_secure()
    ):
        weiter = ""
    if request.method == "POST":
        form = TerminForm(request.POST)
        if form.is_valid():
            termin = form.save(commit=False)
            termin._audit_user = request.user
            termin.save()
            return redirect(weiter) if weiter else redirect("termin_view", id=termin.id)
    else:
        form = TerminForm()
    context = {
        "form": form,
        "seitentitel": "Neuer Termin",
        "next": weiter,
    }
    return render(request, "termin_edit.html", context)


# --- Kalender-Monatsansicht ---

_MONATSNAMEN = [
    "", "Januar", "Februar", "März", "April", "Mai", "Juni",
    "Juli", "August", "September", "Oktober", "November", "Dezember",
]
_WOCHENTAGE = ["Mo", "Di", "Mi", "Do", "Fr", "Sa", "So"]
_WOCHENTAGE_LANG = [
    "Montag", "Dienstag", "Mittwoch", "Donnerstag", "Freitag", "Samstag", "Sonntag",
]


@login_required
def termin_kalender(request):
    """Klassische Monatsansicht: 7-Spalten-Raster (Mo–So) mit Terminen und
    Geburtstagen. Monatswechsel per ?jahr=&monat=; htmx-Requests erhalten nur
    das Raster-Partial für einen Swap ohne Full-Reload."""
    today = timezone_now().date()

    # Monat/Jahr robust aus den GET-Parametern lesen; Default = aktueller Monat.
    try:
        jahr = int(request.GET.get("jahr", today.year))
        monat = int(request.GET.get("monat", today.month))
        first = date(jahr, monat, 1)  # ValueError bei ungültigem Monat/Jahr
    except (TypeError, ValueError):
        jahr, monat, first = today.year, today.month, today.replace(day=1)

    cal = _calendar.Calendar(firstweekday=0)  # 0 = Montag
    weeks = cal.monthdatescalendar(jahr, monat)  # Liste von Wochen à 7 date-Objekten
    grid_start, grid_end = weeks[0][0], weeks[-1][-1]

    # Termine, die den sichtbaren Bereich berühren – auch mehrtägige, die vor dem
    # Raster beginnen oder nach ihm enden – auf jeden betroffenen Tag verteilen.
    termine_by_day = {}
    termin_qs = (
        Termin.objects.filter(datum__date__lte=grid_end)
        .filter(
            Q(datum_bis__date__gte=grid_start)
            | Q(datum_bis__isnull=True, datum__date__gte=grid_start)
        )
        .select_related("ort")
        .order_by("datum")
    )
    for t in termin_qs:
        start = t.datum.date()
        ende = t.datum_bis.date() if t.datum_bis else start
        if ende < start:
            ende = start
        mehrtaegig = ende > start
        tag = max(start, grid_start)
        letzter = min(ende, grid_end)
        while tag <= letzter:
            termine_by_day.setdefault(tag, []).append(
                {"t": t, "mehrtaegig": mehrtaegig, "start": tag == start, "ende": tag == ende}
            )
            tag += timedelta(days=1)

    # Geburtstage: nur aktive Mitglieder, deren Geburtsdatum nicht verborgen ist.
    # Nur Monat/Tag zählen (Geburtsjahr irrelevant), daher auf die im Raster
    # vorkommenden (Monat, Tag)-Paare abbilden.
    grid_md = {(d.month, d.day) for week in weeks for d in week}
    geb_by_md = {}
    for u in User.objects.filter(is_active=True, geburtstag__isnull=False):
        if "geburtstag" in (u.private_fields or []):
            continue
        md = (u.geburtstag.month, u.geburtstag.day)
        if md in grid_md:
            geb_by_md.setdefault(md, []).append(u)
    for lst in geb_by_md.values():
        lst.sort(key=lambda m: (m.stimme_pos, m.first_name, m.last_name))

    wochen = [
        [
            {
                "tag": d.day,
                "label": f"{_WOCHENTAGE_LANG[d.weekday()]}, {d.day}. {_MONATSNAMEN[d.month]} {d.year}",
                "iso": d.isoformat(),  # Vorbelegung des Datumsfelds im Offcanvas
                "im_monat": d.month == monat,
                "ist_heute": d == today,
                "wochenende": d.weekday() >= 5,
                "termine": termine_by_day.get(d, []),
                "geburtstage": geb_by_md.get((d.month, d.day), []),
            }
            for d in week
        ]
        for week in weeks
    ]

    prev_first = (first - timedelta(days=1)).replace(day=1)
    next_first = (weeks[-1][-1] + timedelta(days=1)).replace(day=1)

    context = {
        "seitentitel": f"Kalender · {_MONATSNAMEN[monat]} {jahr}",
        "wochen": wochen,
        "wochentage": _WOCHENTAGE,
        "monatsname": _MONATSNAMEN[monat],
        "jahr": jahr,
        "prev_jahr": prev_first.year,
        "prev_monat": prev_first.month,
        "next_jahr": next_first.year,
        "next_monat": next_first.month,
        "heute_jahr": today.year,
        "heute_monat": today.month,
        "ist_aktueller_monat": (jahr == today.year and monat == today.month),
        # Anlege-Formular im Tages-Offcanvas – nur für Berechtigte; das Raster-
        # Partial braucht es nicht, es lebt im Panel außerhalb von #cal-view.
        "add_form": TerminForm() if request.user.has_perm("intern.add_termin") else None,
    }
    if request.headers.get("HX-Request"):
        return render(request, "_kalender_grid.html", context)
    return render(request, "kalender.html", context)